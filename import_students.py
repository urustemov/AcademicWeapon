"""Validate and synchronize the current Academic Probation workbook.

The command is dry-run by default and never prints student names, IDs, or
emails. Use --apply only with the separate non-demo database configured for the
protected application environment.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
import re

from openpyxl import load_workbook

import db


DEFAULT_SHEET = 'Academic Probation (62)'
SOURCE_NAME = 'academic_probation_2026_2027'
REQUIRED_HEADERS = {
    'name': 'Student Name',
    'student_code': 'Student ID',
    'program': 'Program',
    'year': 'Year of Standing',
    'email': 'Student E-mail',
    'gpa': 'CGPA',
    'advisor': "Advisor's Name, Surname",
}
YEAR_MAP = {'freshman': 1, 'sophomore': 2, 'junior': 3, 'senior': 4}


@dataclass(frozen=True)
class StudentRecord:
    student_code: str
    name: str
    email: str
    program: str
    year: int
    gpa: float
    advisor: str
    avatar: str


class ImportValidationError(ValueError):
    pass


def text(value) -> str:
    return str(value or '').strip()


def normalize_student_id(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text(value)


def initials(name: str) -> str:
    result = ''.join(part[:1] for part in name.split()[:2]).upper()
    return result or 'ST'


def parse_year(value) -> int | None:
    if isinstance(value, (int, float)) and int(value) == value and 1 <= int(value) <= 4:
        return int(value)
    normalized = text(value).lower()
    if normalized in YEAR_MAP:
        return YEAR_MAP[normalized]
    match = re.search(r'\b([1-4])\b', normalized)
    return int(match.group(1)) if match else None


def read_students(path: Path, sheet_name: str = DEFAULT_SHEET) -> list[StudentRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ImportValidationError(f'Required worksheet is missing: {sheet_name}')
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [text(value) for value in next(rows)]
        missing_headers = [header for header in REQUIRED_HEADERS.values() if header not in headers]
        if missing_headers:
            raise ImportValidationError('Missing required columns: ' + ', '.join(missing_headers))
        columns = {key: headers.index(header) for key, header in REQUIRED_HEADERS.items()}

        records: list[StudentRecord] = []
        errors: list[str] = []
        for row_number, row in enumerate(rows, start=2):
            if not any(text(value) for value in row):
                continue
            student_code = normalize_student_id(row[columns['student_code']])
            name = text(row[columns['name']])
            email = text(row[columns['email']]).lower()
            program = text(row[columns['program']])
            year = parse_year(row[columns['year']])
            advisor = db.canonical_advisor_name(text(row[columns['advisor']]))
            try:
                gpa = float(row[columns['gpa']])
            except (TypeError, ValueError):
                gpa = -1

            row_errors = []
            if not student_code:
                row_errors.append('missing Student ID')
            if not name:
                row_errors.append('missing student name')
            if not re.fullmatch(r'[^\s@]+@auca\.kg', email, flags=re.IGNORECASE):
                row_errors.append('email is missing or is not an @auca.kg address')
            if not program:
                row_errors.append('missing program')
            if year is None:
                row_errors.append('invalid year of standing')
            if not 0 <= gpa <= 4:
                row_errors.append('CGPA must be between 0 and 4')
            if not advisor:
                row_errors.append('missing advisor name')
            if row_errors:
                errors.append(f'row {row_number}: ' + '; '.join(row_errors))
                continue

            records.append(StudentRecord(
                student_code=student_code, name=name, email=email, program=program,
                year=year, gpa=gpa, advisor=advisor, avatar=initials(name),
            ))

        duplicate_ids = len(records) - len({record.student_code for record in records})
        duplicate_emails = len(records) - len({record.email for record in records})
        if duplicate_ids:
            errors.append(f'{duplicate_ids} duplicate Student ID value(s)')
        if duplicate_emails:
            errors.append(f'{duplicate_emails} duplicate email value(s)')
        if errors:
            preview = '\n'.join(errors[:20])
            suffix = f'\n...and {len(errors) - 20} more issue(s)' if len(errors) > 20 else ''
            raise ImportValidationError(preview + suffix)
        return records
    finally:
        workbook.close()


def synchronize(records: list[StudentRecord], database_path: Path, deactivate_missing: bool) -> dict[str, int]:
    db.DB_PATH = database_path
    db.DEMO_MODE = False
    db.seed_database()
    now = datetime.now(timezone.utc).isoformat(timespec='seconds')
    created = updated = users_created = deactivated = 0

    with db.connect() as con:
        for record in records:
            existing = con.execute('SELECT * FROM students WHERE student_code=?', (record.student_code,)).fetchone()
            email_owner = con.execute(
                'SELECT id,student_code FROM students WHERE LOWER(email)=? AND student_code<>?',
                (record.email, record.student_code),
            ).fetchone()
            if email_owner:
                raise ImportValidationError('A workbook email conflicts with another student already in the database.')

            if existing:
                student_id = existing['id']
                con.execute('''UPDATE students
                    SET name=?,email=?,program=?,year=?,gpa=?,advisor=?,active=1,source=?,synced_at=?
                    WHERE id=?''', (
                    record.name, record.email, record.program, record.year, record.gpa, record.advisor,
                    SOURCE_NAME, now, student_id,
                ))
                updated += 1
            else:
                cursor = con.execute('''INSERT INTO students
                    (student_code,name,email,program,year,gpa,credit_load,advisor,status,attention,avatar,active,source,synced_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                    record.student_code, record.name, record.email, record.program,
                    record.year, record.gpa, 0, record.advisor, 'Initial Assessment',
                    'Regular', record.avatar, 1, SOURCE_NAME, now,
                ))
                student_id = cursor.lastrowid
                created += 1

            user_by_student = con.execute('SELECT * FROM users WHERE student_id=?', (student_id,)).fetchone()
            user_by_email = con.execute('SELECT * FROM users WHERE LOWER(email)=?', (record.email,)).fetchone()
            if user_by_student and user_by_email and user_by_student['id'] != user_by_email['id']:
                raise ImportValidationError('An imported email belongs to a different existing login account.')
            user = user_by_student or user_by_email
            if user:
                if user['role'] != 'student' or (user['student_id'] not in (None, student_id)):
                    raise ImportValidationError('An imported student conflicts with an existing non-student account.')
                con.execute('UPDATE users SET email=?,display_name=?,student_id=?,active=1 WHERE id=?',
                            (record.email, record.name, student_id, user['id']))
            else:
                con.execute('''INSERT INTO users
                    (email,display_name,role,student_id,password_hash,google_sub,active)
                    VALUES (?,?,?,?,?,?,1)''',
                    (record.email, record.name, 'student', student_id, '!google-only!', None))
                users_created += 1

        if deactivate_missing:
            codes = [record.student_code for record in records]
            placeholders = ','.join('?' for _ in codes)
            cursor = con.execute(
                f'''UPDATE students SET active=0
                    WHERE source=? AND student_code NOT IN ({placeholders}) AND active=1''',
                (SOURCE_NAME, *codes),
            )
            deactivated = cursor.rowcount
            con.execute('''UPDATE users SET active=0
                WHERE role='student' AND student_id IN (
                    SELECT id FROM students WHERE source=? AND active=0
                )''', (SOURCE_NAME,))

    return {
        'records': len(records), 'students_created': created,
        'students_updated': updated, 'login_accounts_created': users_created,
        'students_deactivated': deactivated,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description='Validate or synchronize the Academic Probation student list.')
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--sheet', default=DEFAULT_SHEET)
    parser.add_argument('--db', type=Path, default=db.BASE_DIR / 'academic_weapon.sqlite3')
    parser.add_argument('--apply', action='store_true', help='Write validated records to the selected database.')
    parser.add_argument('--deactivate-missing', action='store_true', help='Deactivate previously imported students absent from this workbook.')
    args = parser.parse_args()

    records = read_students(args.workbook, args.sheet)
    print(f'Validation passed: {len(records)} student records; no duplicate IDs or emails.')
    print('Only the minimum access fields will be synchronized; phone numbers, comments, contact history, and counseling fields are excluded.')
    if not args.apply:
        print('Dry run only: the database was not changed.')
        return 0

    result = synchronize(records, args.db.resolve(), args.deactivate_missing)
    print('Synchronization completed.')
    for key, value in result.items():
        print(f'{key}: {value}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
